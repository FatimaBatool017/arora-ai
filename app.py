import io
import json
import mimetypes
import time

import pandas as pd
import streamlit as st
try:
    from google import genai
    from google.genai import types
except ImportError:
    # Safety net: install the package at runtime if the host skipped requirements.txt
    import importlib
    import subprocess
    import sys

    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", "google-genai>=1.0"])
    importlib.invalidate_caches()
    from google import genai
    from google.genai import types

MODEL = "gemini-2.5-flash"          # change here if Google renames the Flash model
IMAGE_MODEL = "gemini-2.5-flash-image"
BIG_FILE = 15 * 1024 * 1024         # above this, files go through the Gemini Files API

st.set_page_config(page_title="StudyBuddy AI", page_icon="🎓", layout="wide")
st.markdown(
    """<style>
    .stApp{background:#F7FBFF} h1,h2,h3{color:#14506E}
    div.stButton>button,div.stFormSubmitButton>button{background:#3AA6E0;color:#fff;border:0;border-radius:10px}
    .token{background:linear-gradient(135deg,#FFD54F,#FF8A65);padding:28px;border-radius:20px;text-align:center;
           font-size:28px;font-weight:700;color:#4B2E00;margin:12px 0}
    </style>""",
    unsafe_allow_html=True,
)

# ---------------------------------------------------------------- API client
key = ""
try:
    key = st.secrets.get("GEMINI_API_KEY", "")
except Exception:
    pass
if not key:
    key = st.sidebar.text_input("🔑 Gemini API key", type="password", help="Free at aistudio.google.com")
if not key:
    st.title("🎓 StudyBuddy AI")
    st.info("Add your free Gemini API key in the sidebar (or in Streamlit secrets) to start.")
    st.stop()
client = genai.Client(api_key=key)


# ---------------------------------------------------------------- helpers
def read_file(f):
    """Turn an uploaded file into {'name', 'text'} or {'name', 'bytes', 'mime'}."""
    name, data = f.name, f.getvalue()
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    mime = f.type or mimetypes.guess_type(name)[0] or "application/octet-stream"
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            for r in t.rows:
                text += "\n" + " | ".join(c.text for c in r.cells)
        return {"name": name, "text": text}
    if ext == "pptx":
        from pptx import Presentation
        text = ""
        for i, s in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            text += f"\n--- Slide {i} ---\n"
            for sh in s.shapes:
                if sh.has_text_frame:
                    text += sh.text_frame.text + "\n"
        return {"name": name, "text": text}
    if ext in ("xlsx", "xls"):
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
        return {"name": name, "text": "\n\n".join(f"Sheet {k}:\n{v.to_csv(index=False)}" for k, v in sheets.items())}
    if mime.startswith(("application/pdf", "image/", "audio/", "video/")):
        return {"name": name, "bytes": data, "mime": mime}
    try:
        return {"name": name, "text": data.decode("utf-8")}
    except UnicodeDecodeError:
        return None


def to_parts(src):
    if "text" in src:
        return [src["text"]]
    if len(src["bytes"]) < BIG_FILE:
        return [types.Part.from_bytes(data=src["bytes"], mime_type=src["mime"])]
    if "_up" not in src:
        up = client.files.upload(file=io.BytesIO(src["bytes"]), config=types.UploadFileConfig(mime_type=src["mime"]))
        while up.state and up.state.name == "PROCESSING":
            time.sleep(3)
            up = client.files.get(name=up.name)
        src["_up"] = up
    return [src["_up"]]


def ask(prompt, src=None, as_json=False):
    contents = (to_parts(src) if src else []) + [prompt]
    cfg = types.GenerateContentConfig(response_mime_type="application/json") if as_json else None
    for attempt in range(3):
        try:
            return client.models.generate_content(model=MODEL, contents=contents, config=cfg).text
        except Exception as e:
            if "429" in str(e) and attempt < 2:      # free-tier rate limit: wait and retry
                time.sleep(25)
                continue
            raise


def make_image(prompt):
    r = client.models.generate_content(
        model=IMAGE_MODEL, contents=prompt,
        config=types.GenerateContentConfig(response_modalities=["TEXT", "IMAGE"]))
    for p in r.candidates[0].content.parts:
        if getattr(p, "inline_data", None) and p.inline_data.data:
            return p.inline_data.data
    return None


def speak(text):
    from gtts import gTTS
    buf = io.BytesIO()
    gTTS(text, lang="en").write_to_fp(buf)
    return buf.getvalue()


def load_once(slot, f, label=""):
    """Load a newly uploaded/recorded file as the current material, once per new file."""
    if f and st.session_state.get("seen_" + slot) != (f.name, f.size):
        st.session_state["seen_" + slot] = (f.name, f.size)
        s = read_file(f)
        if s:
            s["name"] = label + f.name
            st.session_state.src = s
            st.rerun()
        else:
            st.error("Sorry, I can't read this file type.")


def need_src():
    src = st.session_state.get("src")
    if not src:
        st.warning("Upload a file or record audio first (see the Explain and Lectures tabs).")
    return src


def section(key, label, prompt, src=None):
    if st.button(label, key="b_" + key):
        with st.spinner("Working on it..."):
            try:
                st.session_state[key] = ask(prompt, src)
            except Exception as e:
                st.error(f"AI error: {e}")
    out = st.session_state.get(key)
    if out:
        st.markdown(out)
        st.download_button("⬇️ Download", out, f"{key}.md", key="d_" + key)


def parse_json(t):
    t = t.strip().removeprefix("```json").removeprefix("```").removesuffix("```").strip()
    return json.loads(t)


def make_mcqs(src, topic, total=100, per=25):
    focus = ["key facts and definitions", "concepts and understanding", "real-life application", "tricky distinctions and analysis"]
    qs, seen, bar = [], set(), st.progress(0, text="Writing questions...")
    tries = 0
    while len(qs) < total and tries < 10:
        avoid = "; ".join(q["question"][:50] for q in qs[-40:])
        prompt = (f"Create {per} NEW multiple-choice exam questions with exactly 4 options each"
                  f"{'' if src else ' about: ' + topic}. Focus: {focus[tries % 4]}. Do NOT repeat: {avoid}. "
                  'Return ONLY a JSON list of objects: {"question": str, "options": [4 strings], '
                  '"answer": index 0-3 of the correct option, "explanation": short str}.')
        try:
            for q in parse_json(ask(prompt, src, as_json=True)):
                if (len(q.get("options", [])) == 4 and isinstance(q.get("answer"), int) and 0 <= q["answer"] < 4
                        and q["question"] not in seen):
                    seen.add(q["question"])
                    qs.append(q)
        except Exception as e:
            st.caption(f"Retrying a batch ({e.__class__.__name__})...")
        tries += 1
        bar.progress(min(len(qs) / total, 1.0), text=f"{min(len(qs), total)} / {total} questions")
    bar.empty()
    return qs[:total]


def celebrate(pct):
    st.balloons()
    st.markdown(
        """<style>@keyframes flash{0%{opacity:.95}100%{opacity:0}}
        .flash{position:fixed;inset:0;background:radial-gradient(circle,#fffbd0,#ffb347 45%,#ff5e7e);
        animation:flash 2.8s ease-out forwards;pointer-events:none;z-index:99999}</style><div class="flash"></div>""",
        unsafe_allow_html=True)
    tier = ("🏆 GOLD TOKEN: Outstanding!" if pct >= 90 else "🥈 SILVER TOKEN: Great job!" if pct >= 70
            else "🥉 BRONZE TOKEN: Good effort!" if pct >= 50 else "🌱 GROWTH TOKEN: Every try makes you stronger!")
    st.markdown(f'<div class="token">{tier}</div>', unsafe_allow_html=True)


# ---------------------------------------------------------------- sidebar
st.sidebar.title("🎓 StudyBuddy AI")
cur = st.session_state.get("src")
st.sidebar.write(f"**Current material:** {cur['name'] if cur else 'none yet'}")
if cur and st.sidebar.button("Clear material"):
    st.session_state.pop("src")
    st.rerun()
st.sidebar.caption("Free Gemini keys have daily and per-minute limits. If you see a quota error, wait a minute and retry.")

st.title("🎓 StudyBuddy AI")
tabs = st.tabs(["📂 Explain", "🎙️ Lectures", "📝 Notes", "🗺️ Roadmap", "🎨 Visuals", "✅ Exam"])

# ---------------------------------------------------------------- Explain
with tabs[0]:
    f = st.file_uploader("Upload any file (PDF, Word, PowerPoint, Excel, text, image, audio, video)")
    load_once("explain", f)
    style = st.radio("Explanation style", ["Like I'm 10 (very simple)", "Simple and clear", "Detailed but easy"], horizontal=True)
    src = need_src()
    if src:
        section("explain", "✨ Explain it", (
            f"Explain everything in this material in very easy language. Style: {style}. Structure: "
            "1) the big idea in one line, 2) step-by-step explanation with real-life examples and analogies, "
            "3) a glossary of key terms, 4) common mistakes, 5) a short recap. Use light emojis."), src)

# ---------------------------------------------------------------- Lectures
with tabs[1]:
    st.write("Record a lecture or upload an audio file. Long recordings work best as uploaded files.")
    rec = st.audio_input("🎤 Record")
    up = st.file_uploader("...or upload audio/video", type=["mp3", "wav", "m4a", "aac", "ogg", "flac", "mp4", "mov", "webm"], key="aud")
    load_once("rec", rec, "Lecture: ")
    load_once("aud", up, "Lecture: ")
    src = need_src()
    if src:
        section("lecture", "📌 Summarise this lecture", (
            "Listen to/read this lecture. Give: 1) a 3-line TL;DR, 2) the most important points as clear bullets, "
            "3) key terms with simple meanings, 4) formulas/dates/names to remember, 5) 5 quick self-test questions. "
            "Keep it easy and exam-focused."), src)

# ---------------------------------------------------------------- Notes
with tabs[2]:
    src = need_src()
    if src:
        section("notes", "📝 Make friendly notes", (
            "Turn this material into clean, friendly study notes: short headings, bullets, bold key terms, "
            "simple examples, a mind-map style outline, and a final cheat-sheet. Write in easy words."), src)

# ---------------------------------------------------------------- Roadmap
with tabs[3]:
    field = st.text_input("Field of study", placeholder="e.g. Data Science, Medicine, Graphic Design")
    c1, c2, c3 = st.columns(3)
    level = c1.selectbox("Your level", ["Beginner", "Intermediate", "Advanced", "PhD"])
    amt = c2.number_input("Time to learn", 1, 500, 6)
    unit = c3.selectbox("Unit", ["weeks", "months", "years"], index=1)
    if field.strip():
        section("roadmap", "🗺️ Create my roadmap", (
            f"Create a practical learning roadmap for '{field}'. Learner level: {level}. Time available: {amt} {unit}. "
            "Split into phases with a week-by-week or month-by-month plan, topics, free resources (courses, books, sites), "
            "hands-on projects, milestones, how to measure progress, and tips to stay consistent. "
            "Fit the plan to the time. For PhD level include research directions, key papers, writing and publishing."))

# ---------------------------------------------------------------- Visuals
with tabs[4]:
    st.caption("Image and voice features use free-tier models, which may be limited or unavailable on some accounts.")
    idea = st.text_input("Describe the picture you want (or leave blank to use your material)")
    if st.button("🖼️ Generate image"):
        with st.spinner("Drawing..."):
            try:
                p = idea or ask("Describe in one sentence a simple educational illustration for this material.", need_src())
                img = make_image(f"Colourful, friendly educational illustration: {p}")
                st.image(img) if img else st.warning("The image model returned no picture. Try different words.")
            except Exception as e:
                st.error(f"Image error: {e}")
    st.divider()
    st.subheader("🎬 Cartoon explainer")
    st.caption("A short story of cartoon pictures with a narrator voice that explains your material.")
    if st.button("Create cartoon explainer"):
        src = need_src()
        if src:
            with st.spinner("Writing the story, drawing scenes and recording the voice..."):
                try:
                    scenes = parse_json(ask(
                        "Explain this material as a fun cartoon story in 6 scenes for a beginner. Return ONLY a JSON list of "
                        '{"narration": 2 easy sentences, "image_prompt": cartoon scene description}.', src, as_json=True))[:6]
                    out = []
                    for sc in scenes:
                        out.append((make_image("Cute colourful cartoon style, no text: " + sc["image_prompt"]),
                                    sc["narration"], speak(sc["narration"])))
                    st.session_state.cartoon = out
                except Exception as e:
                    st.error(f"Could not finish the cartoon: {e}")
    for i, (img, text, audio) in enumerate(st.session_state.get("cartoon", []), 1):
        st.markdown(f"**Scene {i}**")
        if img:
            st.image(img, width=420)
        st.write(text)
        st.audio(audio, format="audio/mp3")

# ---------------------------------------------------------------- Exam
with tabs[5]:
    src = st.session_state.get("src")
    topic = "" if src else st.text_input("No material loaded. Enter a topic for the exam")
    total = st.number_input("Number of questions", 10, 150, 100, step=10)
    if st.button("🧪 Generate exam") and (src or topic.strip()):
        for k in [k for k in st.session_state if k.startswith("ans")] + ["qs", "result", "celebrated"]:
            st.session_state.pop(k, None)
        st.session_state.qs = make_mcqs(src, topic, int(total))
    qs = st.session_state.get("qs")
    if qs and "result" not in st.session_state:
        st.success(f"{len(qs)} questions ready. Answer and submit at the bottom.")
        with st.form("exam"):
            for i, q in enumerate(qs):
                st.radio(f"**Q{i + 1}.** {q['question']}", q["options"], index=None, key=f"ans{i}")
            if st.form_submit_button("✅ Submit exam"):
                st.session_state.result = {i: st.session_state.get(f"ans{i}") for i in range(len(qs))}
                st.rerun()
    if qs and "result" in st.session_state:
        res = st.session_state.result
        score = sum(res[i] == q["options"][q["answer"]] for i, q in enumerate(qs))
        pct = round(100 * score / len(qs))
        st.header(f"Your score: {score} / {len(qs)} ({pct}%)")
        if not st.session_state.get("celebrated"):
            celebrate(pct)
            st.session_state.celebrated = True
        with st.expander("📖 Review answers", expanded=True):
            for i, q in enumerate(qs):
                ok = res[i] == q["options"][q["answer"]]
                st.markdown(f"{'✅' if ok else '❌'} **Q{i + 1}. {q['question']}**  \n"
                            f"Your answer: {res[i] or 'not answered'}  \nCorrect: {q['options'][q['answer']]}  \n_{q.get('explanation', '')}_")
        if st.button("🔁 Retake exam"):
            for k in [k for k in st.session_state if k.startswith("ans")] + ["result", "celebrated"]:
                st.session_state.pop(k, None)
            st.rerun()
